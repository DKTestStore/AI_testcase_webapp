
import json
import os
from typing import Any, Dict, List, Optional
import pandas as pd
from openpyxl import load_workbook

# --------------------------------------------------------------------------- #
# 1. COLUMN NAME MAPPING (edit here if your Excel headers ever change)
# --------------------------------------------------------------------------- #
COLUMN_MAP = {
    "id": "Test Case ID",
    "name": "Test Case Name",
    "description": "Test Case Description",
    "precondition": "Precondition",
    "step_description": "Step Description",
    "step_expected": "Expected Result",
}

# --------------------------------------------------------------------------- #
# 2. READ EXCEL -> LIST OF TEST-CASE DICTS
# --------------------------------------------------------------------------- #
def read_test_cases_from_excel(
    filepath: str,
    sheet_name: Optional[str] = 0,
    column_map: Dict[str, str] = COLUMN_MAP,
) -> List[Dict[str, Any]]:
    """
    Read an Excel sheet where test-case-level columns (ID/Name/Description/
    Precondition) are filled only on each test case's first row, and every
    row contributes one Step Description / Expected Result pair.

    Returns a list of dicts:
        [
            {
                "id": "TC[001]",
                "name": "...",
                "description": "...",
                "precondition": "...",
                "steps": [{"description": "...", "expected": "..."}, ...],
            },
            ...
        ]
    """
    df = pd.read_excel(filepath, sheet_name=sheet_name, dtype=str)
    df.columns = [str(c).strip() for c in df.columns]

    # Forward-fill the test-case-level columns so every step row knows which
    # test case it belongs to (handles merged-cell exports from Excel).
    fill_cols = [column_map["id"], column_map["name"], column_map["description"], column_map["precondition"]]
    for col in fill_cols:
        if col in df.columns:
            df[col] = df[col].ffill()

    test_cases: List[Dict[str, Any]] = []
    current_id = None
    current_tc: Optional[Dict[str, Any]] = None

    for _, row in df.iterrows():
        tc_id = _clean(row.get(column_map["id"]))
        step_desc = _clean(row.get(column_map["step_description"]))
        step_expected = _clean(row.get(column_map["step_expected"]))

        # Skip fully blank rows
        if not tc_id and not step_desc and not step_expected:
            continue

        if tc_id and tc_id != current_id:
            # Start a new test case
            if current_tc:
                test_cases.append(current_tc)
            current_id = tc_id
            current_tc = {
                "id": tc_id,
                "name": _clean(row.get(column_map["name"])) or "",
                "description": _clean(row.get(column_map["description"])) or "",
                "precondition": _clean(row.get(column_map["precondition"])) or "",
                "steps": [],
            }

        if current_tc is not None and (step_desc or step_expected):
            current_tc["steps"].append(
                {"description": step_desc or "", "expected": step_expected or ""}
            )

    if current_tc:
        test_cases.append(current_tc)

    return test_cases

def _clean(value: Any) -> Optional[str]:
    """Normalize a cell value: strip whitespace, treat NaN/'nan' as None."""
    if value is None:
        return None
    if isinstance(value, float) and pd.isna(value):
        return None
    text = str(value).strip()
    if text.lower() in ("", "nan", "none"):
        return None
    return text



def save_qtest_results_to_excel(excel_path, payloads, results):
    """Append qTest metadata (name + ID + PID) to a new worksheet."""

    if not os.path.exists(excel_path):
        raise FileNotFoundError(f"Excel file not found: {excel_path}")

    workbook = load_workbook(excel_path)
    sheet_name = "qTest Results"

    if sheet_name in workbook.sheetnames:
        del workbook[sheet_name]

    ws = workbook.create_sheet(title=sheet_name)

    # Updated Header
    ws.append([
        "Test Case Name",
        "qTest ID",
        "Test Case ID",
        "Status Code",
        "Response JSON"
    ])

    for payload, result in zip(payloads, results):

        response = result.get("response") or {}

        # Extract values
        qtest_id = response.get("id")
        pid = response.get("pid")

        # Handle nested response structures if needed
        if isinstance(response.get("data"), dict):
            qtest_id = qtest_id or response["data"].get("id")
            pid = pid or response["data"].get("pid")

        ws.append([
            payload.get("name", ""),
            qtest_id or "",
            pid or "",
            result.get("status_code", ""),
            json.dumps(response, ensure_ascii=False)
        ])

    workbook.save(excel_path)

    print(
        f"Saved qTest result metadata to '{excel_path}' "
        f"-> sheet '{sheet_name}'"
    )



# --------------------------------------------------------------------------- #
# 3. BUILD A SINGLE qTest PAYLOAD FROM ONE TEST-CASE DICT
# --------------------------------------------------------------------------- #
def build_payload(test_case: Dict[str, Any], parent_id: Optional[int] = None) -> Dict[str, Any]:
    """
    Convert one parsed test-case dict (from read_test_cases_from_excel) into
    the qTest 'Create Test Case' JSON payload format.
    """
    payload: Dict[str, Any] = {
        "name": test_case.get("name", ""),
        "properties": [],
        "description": test_case.get("description", ""),
        "precondition": test_case.get("precondition", ""),
        "test_steps": [
            {"description": step["description"], "expected": step["expected"]}
            for step in test_case.get("steps", [])
        ],
    }
    if parent_id is not None:
        payload["parent_id"] = parent_id
    return payload


# --------------------------------------------------------------------------- #
# 4. BUILD BULK PAYLOADS (LIST) FROM ALL TEST CASES IN THE EXCEL FILE
# --------------------------------------------------------------------------- #
def build_bulk_payloads(
    test_cases: List[Dict[str, Any]], parent_id: Optional[int] = None
) -> List[Dict[str, Any]]:
    """Convert a list of parsed test-case dicts into a list of qTest payloads."""
    return [build_payload(tc, parent_id=parent_id) for tc in test_cases]



# Build Payload for Link Test case and requriment
def build_payload_for_linking_requirement_and_testcase(
    testcase_ids: List[int],
    requirement_id: int
) -> Dict[str, Any]:

    return {
        "requirement_id": int(requirement_id),
        "testcase_ids": testcase_ids
    }


# Build Payload for Create Module or Folder
def build_payload_for_create_module(payload) -> Dict[str, Any]:
    return {
        "name": payload,
        "description": ""
    }