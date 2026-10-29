
import csv
import re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


def read_test_case_generator_prompt_file(filepath):
    """
    Reads the system prompt from an external Markdown or Text file.
    """
    try:
        with open(filepath, "r", encoding="utf-8") as file:
            return file.read()
    except FileNotFoundError:
        print(f"Error: Could not find '{filepath}'. Please ensure the file exists.")
        return None


def read_test_cases_reviewer_prompt_file(filepath):
    """
    Reads the reviewer prompt from an external Markdown or Text file.
    """
    try:
        with open(filepath, "r", encoding="utf-8") as file:
            return file.read()
    except FileNotFoundError:
        print(f"Error: Could not find '{filepath}'. Please ensure the file exists.")
        return None


HEADERS = [
    'Test case Id',
    'Test Case Name',
    'Test Case Description',
    'Precondition',
    'Step Description',
    'Expected Result',
]

def _clean(text):
    return '' if text is None else str(text).replace('**', '').strip()

def _extract_field(lines, label):
    for line in lines:
        clean_line = _clean(line)
        if clean_line.lower().startswith(label.lower() + ':'):
            return clean_line.split(':', 1)[1].strip()
    return ''

def _normalize_output(markdown_text):
    text = markdown_text or ''
    replacements = {
        '### **Test Case ID:': 'Test Case ID:',
        '**Test Case ID:**': 'Test Case ID:',
        '### Test Case ID:': 'Test Case ID:',
        'Test Case Id:': 'Test Case ID:',
        'Test case Id:': 'Test Case ID:',
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    return text

def _parse_markdown_table(lines):
    rows = []
    for line in lines:
        raw = line.strip()
        if not raw.startswith('|'):
            continue
        low = raw.lower()
        if ':---' in raw or '---' in raw:
            continue
        if 'step description' in low and 'expected result' in low:
            continue
        if 'test steps' in low and 'expected result' in low:
            continue
        parts = [_clean(p) for p in raw.strip('|').split('|')]
        if len(parts) >= 2 and (parts[0] or parts[1]):
            rows.append((parts[0], parts[1]))
    return rows

def _parse_html_table(block):
    rows = []
    tr_blocks = re.findall(r'<tr[^>]*>(.*?)</tr>', block, flags=re.I | re.S)
    for tr in tr_blocks:
        cells = re.findall(r'<t[dh][^>]*>(.*?)</t[dh]>', tr, flags=re.I | re.S)
        cells = [re.sub(r'<[^>]+>', ' ', c) for c in cells]
        cells = [re.sub(r'\s+', ' ', _clean(c)) for c in cells]
        if len(cells) >= 2:
            if cells[0].lower() in ('step description', 'test steps', 'test step'):
                continue
            if cells[0] or cells[1]:
                rows.append((cells[0], cells[1]))
    return rows

def parse_test_cases(markdown_text):
    text = _normalize_output(markdown_text)
    blocks = text.split('Test Case ID:')
    parsed_rows = []
    for block in blocks[1:]:
        block = block.strip()
        if not block:
            continue
        lines = block.splitlines()
        test_case_id = _clean(lines[0])
        test_case_name = _extract_field(lines, 'Test Case Name')
        test_case_desc = _extract_field(lines, 'Test Case Description')
        precondition = _extract_field(lines, 'Precondition')

        step_rows = _parse_markdown_table(lines)
        if not step_rows:
            step_rows = _parse_html_table(block)

        for step, expected in step_rows:
            parsed_rows.append([
                test_case_id,
                test_case_name,
                test_case_desc,
                precondition,
                step,
                expected,
            ])
    return parsed_rows

def save_to_csv(markdown_text, filename='Manual_Test_Cases.csv'):
    rows = parse_test_cases(markdown_text)
    with open(filename, mode='w', newline='', encoding='utf-8-sig') as file:
        writer = csv.writer(file)
        writer.writerow(HEADERS)
        writer.writerows(rows)
    print(f"Success! Manual test cases have been saved to '{filename}' with {len(rows)} step rows.")


def save_to_excel(markdown_text, filename="Manual_Test_Cases.xlsx"):

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Test Cases"

    headers = [
        'Test Case ID',
        'Test Case Name',
        'Test Case Description',
        'Precondition',
        'Step Description',
        'Expected Result'
    ]

    ws.append(headers)

    for cell in ws[1]:
        cell.font = Font(bold=True)

    text = _normalize_output(markdown_text)
    blocks = text.split("Test Case ID:")

    for block in blocks[1:]:
        lines = block.strip().splitlines()
        tc_id = _clean(lines[0])
        tc_name = _extract_field(lines, "Test Case Name")
        tc_desc = _extract_field(lines, "Test Case Description")
        precondition = _extract_field(lines, "Precondition")

        step_rows = _parse_markdown_table(lines)

        if not step_rows:
            step_rows = _parse_html_table(block)
        is_first_step = True

        for step, expected in step_rows:
            if is_first_step:
                ws.append([
                    tc_id,
                    tc_name,
                    tc_desc,
                    precondition,
                    step,
                    expected
                ])
                is_first_step = False
            else:
                ws.append([
                    "",
                    "",
                    "",
                    "",
                    step,
                    expected
                ])
    wb.save(filename)

    print(f"✅ Success! Manual test cases saved to '{filename}'"
    )
