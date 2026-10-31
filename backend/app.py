"""AI TestCase Generator - FastAPI backend.

Run from the project root (the folder containing backend/ and frontend/):
    uvicorn backend.app:app --reload
then open http://localhost:8000
"""
import io
import json
import logging
import re
import time
import uuid
import zipfile
from datetime import datetime
from logging.handlers import RotatingFileHandler
from pathlib import Path
from typing import Any, Dict, List, Optional

from fastapi import Body, Depends, FastAPI, Request
from fastapi.responses import FileResponse, JSONResponse, RedirectResponse, Response, StreamingResponse
from fastapi.staticfiles import StaticFiles

from . import config  # noqa: F401  (sys.path + .env)
from . import excel_io, history_store, parsing, services, session_store, settings_store
from . import token_usage
from . import prompt_templates as ptpl
from .config import FRONTEND_DIR, GENERATOR_PROMPT_FILE, LOG_DIR, RUNS_DIR, env
from .services import ServiceError

# Rotating server log: backend/logs/app.log, 5 MB x 3 backups
log = logging.getLogger("atg")
if not log.handlers:
    _h = RotatingFileHandler(LOG_DIR / "app.log", maxBytes=5 * 1024 * 1024, backupCount=3, encoding="utf-8")
    _h.setFormatter(logging.Formatter("%(asctime)s %(levelname)s %(name)s: %(message)s"))
    log.addHandler(_h)
    log.setLevel(logging.INFO)
    for _n in ("uvicorn.error", "uvicorn.access"):
        logging.getLogger(_n).addHandler(_h)

app = FastAPI(title="AI TestCase Generator")
COOKIE = "atg_session"
HCOOKIE = "atg_history"   # long-lived, separate from the session cookie; survives Logout / New run
HMAX_AGE = 60 * 60 * 24 * 365
MAX_TEXT = 200_000


@app.exception_handler(ServiceError)
async def _service_error(_: Request, exc: ServiceError):
    return JSONResponse({"ok": False, "error": exc.message}, status_code=exc.status)


@app.exception_handler(Exception)
async def _unexpected(_: Request, exc: Exception):
    log.exception("Unexpected server error")
    return JSONResponse({"ok": False, "error": f"Unexpected server error: {exc}"}, status_code=500)


# --------------------------------------------------------------------------- #
# Session helpers (no username/password - the Welcome button starts a session)
# --------------------------------------------------------------------------- #
def current(request: Request) -> Dict[str, Any]:
    sess = session_store.get(request.cookies.get(COOKIE))
    if not sess:
        raise ServiceError("Session expired. Please start again from the Welcome page.", 401)
    hid = request.cookies.get(HCOOKIE)
    if history_store.valid(hid) and sess.get("hid") != hid:
        sess["hid"] = hid
    return sess


def _run(sess) -> Dict[str, Any]:
    """The run the step endpoints work on: the single run, or (in Batch mode) the open batch item's run.
    The single run itself is never touched while Batch mode is on."""
    if sess.get("mode") == "batch":
        it = _bitem(sess.get("batch"), (sess.get("batch") or {}).get("active"))
        if it:
            return it["run"]
    return sess["run"]


def _bitem(batch, key):
    return next((i for i in (batch or {}).get("items", []) if i["key"] == key), None) if key else None


def _save(sess, history=False):
    """Persist the session. history=True also upserts the run into the rolling Run history
    (only on real saves: generate / review save / import / Excel / qTest upload)."""
    if history and sess.get("hid"):
        try:
            if sess.get("mode") == "batch" and sess.get("batch"):
                history_store.record_batch(sess["hid"], sess["batch"])
            else:
                history_store.record(sess["hid"], sess["run"])
        except Exception:
            log.exception("history record failed")
    session_store.save(sess["sid"])


def _tok(sess, action, usage, run=None, item=None):
    """Token Usage ledger (additive). Never raises."""
    try:
        run = run if run is not None else _run(sess)
        batch = sess.get("batch") if sess.get("mode") == "batch" else None
        if batch and item is None:
            item = next((i for i in batch.get("items", []) if i.get("run") is run), None)
        mode = "batch" if batch and item is not None else (run.get("source") or "local")
        token_usage.record(action=action, usage=usage, run=run, mode=mode,
                           batch=batch if mode == "batch" else None, item=item)
    except Exception:
        log.exception("token usage hook failed (ignored)")


def _steps(run) -> Dict[str, bool]:
    """Which steps have data (used to enable Next / sidebar progress)."""
    return {
        "input": bool(run.get("feature_text")) or bool(run.get("imported")),
        "generate": bool(run.get("test_cases")),
        "review": bool(run.get("test_cases")),  # Bug 01: Excel download no longer required
        "qtest": bool(run.get("upload")),
    }


def _public_run(run) -> Dict[str, Any]:
    r = dict(run)
    r["steps"] = _steps(run)
    if not r.get("ac_list") and (run.get("feature_text") or run.get("jira")):
        r["ac_list"] = _ac_list(run)  # coverage check needs the AC list even before Generate (e.g. imported cases)
    r["excel_name"] = Path(run["excel_path"]).name if run.get("excel_path") else ""
    r.pop("excel_path", None)
    return r


# --------------------------------------------------------------------------- #
# Pages
# --------------------------------------------------------------------------- #
@app.get("/")
def root(request: Request):
    if session_store.get(request.cookies.get(COOKIE)):
        return RedirectResponse("/app")
    return FileResponse(FRONTEND_DIR / "welcome.html")


@app.get("/welcome")
def welcome():
    return FileResponse(FRONTEND_DIR / "welcome.html")


@app.get("/app")
def app_page(request: Request):
    if not session_store.get(request.cookies.get(COOKIE)):
        return RedirectResponse("/welcome")
    return FileResponse(FRONTEND_DIR / "app.html")


app.mount("/static", StaticFiles(directory=str(FRONTEND_DIR)), name="static")


# --------------------------------------------------------------------------- #
# Session
# --------------------------------------------------------------------------- #
@app.post("/api/session/start")
def start_session(request: Request):
    sess = session_store.get(request.cookies.get(COOKIE))
    sid = sess["sid"] if sess else session_store.create()
    resp = JSONResponse({"ok": True, "redirect": "/app#setup"})  # Enh 02: setup checklist first
    resp.set_cookie(COOKIE, sid, httponly=True, samesite="lax", max_age=60 * 60 * 24 * 7)
    hid = request.cookies.get(HCOOKIE)
    if not history_store.valid(hid):
        hid = history_store.new_hid()
    resp.set_cookie(HCOOKIE, hid, httponly=True, samesite="lax", max_age=HMAX_AGE)  # refresh expiry
    log.info("session start")
    return resp


@app.post("/api/session/logout")
def logout(request: Request):
    sid = request.cookies.get(COOKIE)
    if sid and session_store.get(sid):
        session_store.delete(sid)
    resp = JSONResponse({"ok": True, "redirect": "/welcome"})
    resp.delete_cookie(COOKIE)  # history cookie is intentionally kept
    log.info("session logout")
    return resp


@app.get("/api/state")
def state(sess=Depends(current)):
    return {"ok": True, "run": _public_run(sess["run"]), "mode": sess.get("mode") or "single"}


@app.post("/api/run/new")
def new_run(sess=Depends(current)):
    _save(sess, history=True)  # keep the finished run in history; history itself is never cleared
    run = session_store.reset_run(sess["sid"])
    return {"ok": True, "run": _public_run(run)}


# --------------------------------------------------------------------------- #
# Step 1 - Input
# --------------------------------------------------------------------------- #
@app.post("/api/input/jira")
def input_jira(data: Dict[str, Any] = Body(...), sess=Depends(current)):
    info = services.fetch_jira(str(data.get("story_id", "")))
    run = _run(sess)
    run.update({"source": "jira", "story_id": info["issue_key"], "jira": info,
                "feature_text": info["feature_text"], "local_filename": ""})
    _save(sess)
    return {"ok": True, "jira": info, "run": _public_run(run)}


@app.post("/api/input/local")
def input_local(data: Dict[str, Any] = Body(...), sess=Depends(current)):
    text = str(data.get("text", "")).strip()
    if not text:
        raise ServiceError("Feature description is empty.")
    if len(text) > MAX_TEXT:
        raise ServiceError(f"Feature description is too long (max {MAX_TEXT:,} characters).")
    run = _run(sess)
    run.update({"source": "local", "story_id": str(data.get("story_id", "")).strip(), "jira": None,
                "feature_text": text, "local_filename": str(data.get("filename", ""))[:200]})
    _save(sess)
    return {"ok": True, "run": _public_run(run)}


@app.post("/api/input/feature_text")
def update_feature_text(data: Dict[str, Any] = Body(...), sess=Depends(current)):
    """Lets the user edit the 'Will be sent to generator' text for Jira input too."""
    run = _run(sess)
    run["feature_text"] = str(data.get("text", ""))[:MAX_TEXT]
    _save(sess)
    return {"ok": True, "run": _public_run(run)}


# --------------------------------------------------------------------------- #
# Step 2 - Generate
# --------------------------------------------------------------------------- #
@app.get("/api/prompt/default")
def default_prompt(sess=Depends(current)):
    return {"ok": True, "filename": config.GENERATOR_PROMPT_FILE.name, "text": services.read_default_prompt(),
            "modified": bool(ptpl.default_override())}


@app.put("/api/prompt/default")
def put_default_prompt(data: Dict[str, Any] = Body(...), sess=Depends(current)):
    text = str(data.get("text") or "")
    if not text.strip():
        raise ServiceError("The default prompt cannot be empty.")
    if len(text) > MAX_TEXT:
        raise ServiceError("Prompt is too long.")
    ptpl.update_default({"text": text})   # stored as the edited Default template (original file untouched)
    log.info("default prompt updated (%d chars)", len(text))
    return {"ok": True, "filename": GENERATOR_PROMPT_FILE.name}


@app.get("/api/usage")
def usage(sess=Depends(current)):
    return {"ok": True, "usage": services.read_gemini_usage(), "last": _run(sess).get("last_usage")}


@app.post("/api/generate")
def generate(data: Dict[str, Any] = Body(...), sess=Depends(current)):
    run = _run(sess)
    model = str(data.get("model", "gemini")).lower()
    mode = "custom" if data.get("prompt_mode") == "custom" else "default"
    custom = str(data.get("custom_prompt", ""))[:50_000]
    run_reviewer = bool(data.get("run_reviewer"))
    run.update({"model": model, "prompt_mode": mode, "custom_prompt": custom, "run_reviewer": run_reviewer})
    _save(sess)

    tpl_ids = _tpl_ids_for(run, data) if mode == "default" else None
    tpl_res = None
    if tpl_ids:   # prompt templates (additive): only when something other than the built-in Default is selected
        tpl_res = _tpl_generate(run, model, tpl_ids)
        result, cases, appended = tpl_res["result"], tpl_res["cases"], False
    else:
        prompt, appended = services.build_prompt(mode, custom)
        result = services.generate(model, run["feature_text"], prompt)
        cases = parsing.markdown_to_test_cases(result["markdown"])
        run["prompts_used"] = [{"id": ptpl.DEFAULT_ID if mode == "default" else "custom",
                                "name": "Default" if mode == "default" else "Custom", "prompt": prompt}]
        if mode == "default" and cases:   # Default template implicitly used -> count it
            ptpl.mark_used([ptpl.DEFAULT_ID])
    if not cases:
        run["markdown"] = result["markdown"]
        _save(sess)
        raise ServiceError("The model responded, but no test cases could be parsed from its output "
                           "(expected the 'Test Case ID: … | Step Description | Expected Result |' template).", 422)
    run["ac_list"] = _ac_list(run)
    parsing.suggest_acs(cases, run["ac_list"])
    cases = parsing.sanitize_cases(cases)
    run.update({"imported": False, "markdown": result["markdown"], "test_cases": cases, "dirty": True, "excel_path": "",
                "review": None, "upload": None, "rtm": [], "last_usage": result["usage"]})
    review_info = None
    if run_reviewer:
        try:
            rv = services.review(model, run["feature_text"], result["markdown"])
            rv_cases = parsing.markdown_to_test_cases(rv["markdown"])
            review_info = {"markdown": rv["markdown"], "test_cases": rv_cases, "usage": rv["usage"]}
        except ServiceError as ex:
            review_info = {"error": ex.message}
        run["review"] = review_info
    _save(sess, history=True)
    if tpl_res:
        for u in tpl_res["usages"]:
            _tok(sess, "Generate", u, run)
    else:
        _tok(sess, "Generate", result["usage"], run)
    if review_info and review_info.get("usage"):
        _tok(sess, "AI review", review_info["usage"], run)
    return {"ok": True, "count": len(cases), "template_appended": appended,
            "usage": result["usage"], "review": review_info, "run": _public_run(run)}


# --------------------------------------------------------------------------- #
# Step 3 - Review
# --------------------------------------------------------------------------- #
def _validate_cases(cases: List[Dict[str, Any]]):
    ids = [c["id"] for c in cases]
    if any(not i for i in ids):
        raise ServiceError("Every test case needs a Test case Id.")
    dup = {i for i in ids if ids.count(i) > 1}
    if dup:
        raise ServiceError("Duplicate Test case Id(s): " + ", ".join(sorted(dup)))


@app.put("/api/review/test_cases")
def save_cases(data: Dict[str, Any] = Body(...), sess=Depends(current)):
    cases = parsing.sanitize_cases(data.get("test_cases"))
    _validate_cases(cases)
    run = _run(sess)
    run.update({"test_cases": cases, "dirty": True})
    _save(sess, history=True)
    return {"ok": True, "run": _public_run(run)}


@app.post("/api/review/regenerate")
def regenerate_one(data: Dict[str, Any] = Body(...), sess=Depends(current)):
    """Regenerate ONE test case. Returns a proposal only; accept / add-as-new / undo happen in the
    browser and are persisted with PUT /api/review/test_cases."""
    run = _run(sess)
    tcs = parsing.sanitize_cases([data.get("test_case") or {}])
    if not tcs or not (tcs[0]["name"] or tcs[0]["steps"]):
        raise ServiceError("No test case to regenerate.")
    tc = tcs[0]
    instruction = str(data.get("instruction") or "").strip()[:1000]
    ac_list = run.get("ac_list") or _ac_list(run)
    linked = [a for a in ac_list if a["id"] in tc["acs"]]
    res = services.regenerate_one(run.get("story_id", ""), run.get("feature_text", ""), linked, tc, instruction)
    prop = res["test_case"]
    if data.get("keep_title_priority"):
        prop["name"], prop["priority"] = tc["name"], tc["priority"]
    prop["priority"] = prop["priority"] or tc["priority"]
    prop.update({"acs": tc["acs"], "status": "not", "regenerated": True})
    log.info("regenerate %s (%s)", tc.get("id"), "mock" if res["usage"].get("mock") else "gemini")
    _tok(sess, "Regenerate", res["usage"], run)
    return {"ok": True, "proposal": prop, "usage": res["usage"]}


@app.post("/api/review/ai")
def run_reviewer(sess=Depends(current)):
    run = _run(sess)
    md = parsing.test_cases_to_markdown(run["test_cases"])
    rv = services.review(run.get("model", "gemini"), run["feature_text"], md)
    rv_cases = parsing.markdown_to_test_cases(rv["markdown"])
    run["review"] = {"markdown": rv["markdown"], "test_cases": rv_cases, "usage": rv["usage"]}
    _save(sess)
    _tok(sess, "AI review", rv["usage"], run)
    return {"ok": True, "review": run["review"],
            "summary": {"before": len(run["test_cases"]), "after": len(rv_cases)}}


@app.post("/api/review/accept")
def accept_review(sess=Depends(current)):
    run = _run(sess)
    rv = run.get("review") or {}
    if not rv.get("test_cases"):
        raise ServiceError("No reviewed test cases to apply.")
    rv_cases = rv["test_cases"]
    parsing.suggest_acs(rv_cases, run.get("ac_list") or _ac_list(run))
    run.update({"test_cases": parsing.sanitize_cases(rv_cases), "dirty": True, "review": None})
    _save(sess, history=True)
    return {"ok": True, "run": _public_run(run)}


def _write_excel(run) -> Path:
    _validate_cases(run["test_cases"])
    if not run["test_cases"]:
        raise ServiceError("No test cases to save.")
    md = parsing.test_cases_to_markdown(run["test_cases"])
    name = (run.get("story_id") or "LOCAL").replace("/", "_")
    path = RUNS_DIR / run["run_id"] / f"Manual_test_cases_{name}.xlsx"
    services.save_excel(md, path)
    # Review-only columns (AC Tags / Priority / Review Status) + 'AC List' sheet; never read by the qTest payload builder
    excel_io.add_review_columns(path, run["test_cases"], run.get("ac_list"))
    run.update({"excel_path": str(path), "dirty": False})
    return path


@app.post("/api/review/save_excel")
def save_excel(sess=Depends(current)):
    run = _run(sess)
    _write_excel(run)
    _save(sess, history=True)
    return {"ok": True, "run": _public_run(run)}


@app.get("/api/review/excel")
def download_excel(sess=Depends(current)):
    run = _run(sess)
    if not run.get("excel_path") or run.get("dirty") or not Path(run["excel_path"]).exists():
        _write_excel(run)
        _save(sess, history=True)
    p = Path(run["excel_path"])
    return FileResponse(p, filename=p.name,
                        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


# --------------------------------------------------------------------------- #
# Excel upload (Review + Input) and template
# --------------------------------------------------------------------------- #
@app.get("/api/excel/template")
def excel_template(sess=Depends(current)):
    return StreamingResponse(io.BytesIO(excel_io.template_bytes()),
                             media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": 'attachment; filename="Test_Case_Upload_Template.xlsx"'})


def _parse_upload(data: Dict[str, Any]) -> Dict[str, Any]:
    raw = excel_io.decode_upload(str(data.get("filename", "")), str(data.get("data", "")))
    return excel_io.parse_workbook(raw)


@app.post("/api/excel/import/preview")
def excel_import_preview(data: Dict[str, Any] = Body(...), sess=Depends(current)):
    res = _parse_upload(data)
    res.pop("cases")
    return {"ok": True, **res}


@app.post("/api/excel/import")
def excel_import(data: Dict[str, Any] = Body(...), sess=Depends(current)):
    res = _parse_upload(data)
    mode = str(data.get("mode", "replace"))
    if mode not in ("replace", "append", "merge"):
        raise ServiceError("Unknown import mode.")
    run = _run(sess)
    current = run.get("test_cases") or []
    cases = parsing.sanitize_cases(excel_io.apply_mode(current, res["cases"], mode if current else "replace",
                                                       res["needs_renumber"]))
    _validate_cases(cases)
    # AC list: keep the run's ACs (from the input) and add any new tags from the file
    acs = {a["id"]: a for a in (run.get("ac_list") or [])}
    for a in res["ac_list"]:
        acs.setdefault(a["id"], a)
    run.update({"test_cases": cases, "dirty": True, "excel_path": "", "review": None, "upload": None, "rtm": [],
                "imported": True, "import_info": {"filename": str(data.get("filename", ""))[:200], "mode": mode,
                                                  "count": len(res["cases"])},
                "ac_list": sorted(acs.values(), key=lambda a: (len(a["id"]), a["id"]))})
    if not run.get("story_id") and data.get("story_id"):
        run["story_id"] = str(data.get("story_id")).strip()
    _save(sess, history=True)
    return {"ok": True, "count": len(cases), "imported": len(res["cases"]), "warnings": res["warnings"],
            "run": _public_run(run)}


# --------------------------------------------------------------------------- #
# Step 4 - qTest
# --------------------------------------------------------------------------- #
@app.get("/api/qtest/defaults")
def qtest_defaults(sess=Depends(current)):
    return {"ok": True, "project_id": env("QTEST_PROJECT_ID"), "requirement_id": env("QTEST_REQUIREMENT_ID"),
            "parent_module_id": env("QTEST_PARENT_MODULE_ID"), "configured": bool(env("QTEST_BASE_URL") and env("QTEST_TOKEN"))}


@app.post("/api/qtest/preview")
def qtest_preview(data: Dict[str, Any] = Body(default={}), sess=Depends(current)):
    run = _run(sess)
    if run.get("dirty") or not run.get("excel_path") or not Path(run["excel_path"]).exists():
        _write_excel(run)
        _save(sess)
    payloads = services.build_payload_preview(Path(run["excel_path"]), str(data.get("module_id", "")))
    return {"ok": True, "count": len(payloads), "payloads": payloads}


def _ac_list(run) -> List[Dict[str, str]]:
    jira = run.get("jira") or {}
    ac_source = jira.get("acceptance_criteria") or jira.get("scenarios") or run.get("feature_text", "")
    return parsing.split_acceptance_criteria(ac_source)


def _build_rtm(run) -> List[Dict[str, Any]]:
    acs = _ac_list(run)
    run["ac_list"] = acs
    mapping = parsing.map_cases_to_acs(run["test_cases"], acs)
    titles = {a["id"]: a["title"] for a in acs}
    up = run.get("upload") or {}
    by_tc = {i["tc_id"]: i for i in up.get("items", [])}
    rows = []
    for tc in run["test_cases"]:
        it = by_tc.get(tc["id"], {})
        ac = mapping.get(tc["id"], "")
        rows.append({
            "story_id": run.get("story_id") or "LOCAL", "ac_id": ac, "ac_title": titles.get(ac, ""),
            "tc_id": tc["id"], "name": tc["name"],
            "qtest_pid": it.get("qtest_pid") or "", "qtest_id": it.get("qtest_id") or "",
            "upload_status": it.get("status", "Not uploaded"), "error": it.get("error", ""),
            "link_status": "Linked" if it.get("linked") else "Not linked",
            "module": f"{up.get('module_name') or ''} ({up.get('module_id')})" if up.get("module_id") else "",
        })
    return rows


@app.post("/api/qtest/upload")
def qtest_upload(data: Dict[str, Any] = Body(...), sess=Depends(current)):
    run = _run(sess)
    if not run.get("test_cases"):
        raise ServiceError("No test cases. Generate and review first.")
    if run.get("dirty") or not run.get("excel_path"):
        _write_excel(run)
    result = services.qtest_upload(
        Path(run["excel_path"]), run["test_cases"], run.get("story_id", ""),
        str(data.get("project_id", "")), str(data.get("requirement_id", "")),
        str(data.get("parent_module_id", "")), str(data.get("module_id", "")).strip(),
        str(data.get("module_name", "")).strip())
    run["upload"] = result
    run["rtm"] = _build_rtm(run)
    _save(sess, history=True)
    return {"ok": True, "upload": result, "rtm": run["rtm"], "ac_list": run["ac_list"], "run": _public_run(run)}


@app.post("/api/qtest/retry_failed")
def qtest_retry(sess=Depends(current)):
    run = _run(sess)
    if not run.get("upload"):
        raise ServiceError("Nothing uploaded yet.")
    run["upload"] = services.qtest_retry_failed(run["upload"], run["test_cases"])
    run["rtm"] = _build_rtm(run)
    _save(sess, history=True)
    return {"ok": True, "upload": run["upload"], "rtm": run["rtm"], "run": _public_run(run)}


@app.get("/api/qtest/rtm")
def qtest_rtm(sess=Depends(current)):
    run = _run(sess)
    if run.get("test_cases"):
        overrides = {r["tc_id"]: r["ac_id"] for r in run.get("rtm", []) if r.get("ac_overridden")}
        run["rtm"] = _build_rtm(run)
        for r in run["rtm"]:
            if r["tc_id"] in overrides:
                r["ac_id"], r["ac_overridden"] = overrides[r["tc_id"]], True
                r["ac_title"] = next((a["title"] for a in run["ac_list"] if a["id"] == r["ac_id"]), "")
        _save(sess)
    return {"ok": True, "rtm": run.get("rtm", []), "ac_list": run.get("ac_list", []), "upload": run.get("upload")}


@app.put("/api/qtest/rtm/mapping")
def rtm_mapping(data: Dict[str, Any] = Body(...), sess=Depends(current)):
    """Manually override the heuristic TC -> AC mapping."""
    run = _run(sess)
    tc_id, ac_id = str(data.get("tc_id", "")), str(data.get("ac_id", ""))
    for r in run.get("rtm", []):
        if r["tc_id"] == tc_id:
            r["ac_id"], r["ac_overridden"] = ac_id, True
            r["ac_title"] = next((a["title"] for a in run.get("ac_list", []) if a["id"] == ac_id), "")
    _save(sess)
    return {"ok": True, "rtm": run["rtm"]}


@app.get("/api/qtest/rtm/excel")
def rtm_excel(sess=Depends(current)):
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    run = _run(sess)
    rows = run.get("rtm") or []
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "RTM"
    headers = ["Jira Story ID", "Acceptance Criteria", "AC Title", "Generated TC ID", "Test Case Name",
               "qTest TC ID", "qTest Internal ID", "Upload Status", "Link Status", "Module", "Error"]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="4361EE")
    for r in rows:
        ws.append([r["story_id"], r["ac_id"], r["ac_title"], r["tc_id"], r["name"], r["qtest_pid"],
                   r["qtest_id"], r["upload_status"], r["link_status"], r["module"], r.get("error", "")])
    for col, w in zip("ABCDEFGHIJK", (14, 12, 40, 14, 45, 12, 16, 14, 12, 28, 40)):
        ws.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    name = f"RTM_{run.get('story_id') or 'LOCAL'}.xlsx"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{name}"'})


# --------------------------------------------------------------------------- #
# Settings (stored server-side in ai_test_generator/.env, never returned in full)
# --------------------------------------------------------------------------- #
@app.get("/api/settings")
def get_settings(sess=Depends(current)):
    return {"ok": True, "fields": settings_store.get_settings()}


@app.put("/api/settings")
def put_settings(data: Dict[str, Any] = Body(...), sess=Depends(current)):
    changed = settings_store.update_settings(data.get("values") or {})
    return {"ok": True, "changed": changed, "fields": settings_store.get_settings()}


@app.delete("/api/settings/{key}")
def delete_setting(key: str, sess=Depends(current)):
    if not settings_store.clear_setting(key):
        raise ServiceError("Unknown setting.")
    return {"ok": True, "fields": settings_store.get_settings()}


def _prompt_ok() -> bool:
    try:
        return bool(services.read_default_prompt().strip())
    except Exception:
        return False


@app.get("/api/setup/status")
def setup_status(sess=Depends(current)):
    """Enh 02 - what is configured (booleans only, no secret values)."""
    v = lambda k: bool((env(k) or "").strip().strip("\"'"))
    items = [
        {"key": "gemini_key", "label": "Gemini API key", "required": True, "ok": v("GEMINI_API_KEY"),
         "hint": "Required to generate test cases"},
        {"key": "gemini_model", "label": "Gemini model", "required": True, "ok": v("GEMINI_MODEL"),
         "hint": (env("GEMINI_MODEL") or "").strip().strip("\"'") or "Required to generate test cases"},
        {"key": "jira", "label": "Jira", "required": False,
         "ok": v("JIRA_URL") and v("JIRA_EMAIL") and v("JIRA_API_TOKEN"),
         "hint": "Needed only for Jira story input (Local text works without it)"},
        {"key": "qtest", "label": "qTest", "required": False,
         "ok": v("QTEST_BASE_URL") and v("QTEST_TOKEN") and v("QTEST_PROJECT_ID"),
         "hint": "Needed only for qTest Upload (Excel download works without it)"},
        {"key": "default_prompt", "label": "Default prompt", "required": True, "ok": _prompt_ok(),
         "hint": GENERATOR_PROMPT_FILE.name, "link": "View / Edit in Prompt templates →"},
    ]
    return {"ok": True, "items": items}


# --------------------------------------------------------------------------- #
# Run history (rolling last 5 per browser)
# --------------------------------------------------------------------------- #
def _hid(sess):
    hid = sess.get("hid")
    if not history_store.valid(hid):
        raise ServiceError("Run history is not available in this browser (cookies blocked?).")
    return hid


@app.get("/api/history")
def history_list(sess=Depends(current)):
    items = history_store.list_items(sess.get("hid"))
    cur = sess["batch"]["batch_id"] if sess.get("mode") == "batch" and sess.get("batch") else sess["run"]["run_id"]
    return {"ok": True, "items": items, "max": settings_store.get_history_limit(), "current": cur}


@app.get("/api/settings/history_limit")
def get_history_limit(sess=Depends(current)):
    return {"ok": True, "value": settings_store.get_history_limit(), "min": settings_store.HISTORY_LIMIT_MIN,
            "max": settings_store.HISTORY_LIMIT_MAX, "default": settings_store.HISTORY_LIMIT_DEFAULT}


@app.put("/api/settings/history_limit")
def put_history_limit(data: Dict[str, Any] = Body(...), sess=Depends(current)):
    try:
        n = settings_store.set_history_limit(data.get("value"))
    except ValueError as e:
        raise ServiceError(str(e))
    removed = history_store.prune_all(n, sess.get("hid"))
    log.info("history limit set to %d (pruned %d oldest)", n, removed)
    return {"ok": True, "value": n, "pruned": removed}


@app.post("/api/history/{run_id}/open")
def history_open(run_id: str, sess=Depends(current)):
    hid = _hid(sess)
    snap = history_store.load_run(hid, run_id)
    if not snap:
        raise ServiceError("This run is no longer in history.", 404)
    if snap.get("kind") == "batch":
        if (sess.get("batch") or {}).get("batch_id") != run_id:
            _save(sess, history=True)
        sess.update(mode="batch", batch=snap)
        snap["active"] = None
        _save(sess)
        log.info("history open batch %s", snap.get("rid"))
        return {"ok": True, "mode": "batch", "batch": _public_batch(snap), "run": _public_run(sess["run"])}
    if _run(sess)["run_id"] != run_id:
        _save(sess, history=True)  # keep the run we are leaving
    sess["mode"] = "single"
    sess["run"] = snap
    _save(sess)
    log.info("history open %s", snap.get("rid"))
    return {"ok": True, "run": _public_run(snap)}


@app.get("/api/history/{run_id}/excel")
def history_excel(run_id: str, item: str = "", sess=Depends(current)):
    hid = _hid(sess)
    live = (sess.get("batch") or {}).get("batch_id") == run_id
    run = sess["batch"] if live else (_run(sess) if _run(sess)["run_id"] == run_id else history_store.load_run(hid, run_id))
    if run and run.get("kind") == "batch":
        if item:   # one story run inside a batch
            it = next((i for i in run["items"] if i["run"]["run_id"] == item), None)
            if not it or not it["run"].get("test_cases"):
                raise ServiceError("This story run is no longer in history.", 404)
            r = it["run"]
            if not r.get("excel_path") or r.get("dirty") or not Path(r["excel_path"]).exists():
                _write_excel(r)
            p = Path(r["excel_path"])
            return FileResponse(p, filename=p.name,
                                media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        return _batch_zip(run, sess if live else None, None if live else hid)
    if not run:
        raise ServiceError("This run is no longer in history.", 404)
    if not run.get("excel_path") or run.get("dirty") or not Path(run["excel_path"]).exists():
        _write_excel(run)
        if run is _run(sess):
            _save(sess)
        else:
            history_store.save_snapshot(hid, run)
    p = Path(run["excel_path"])
    return FileResponse(p, filename=p.name,
                        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.delete("/api/history/{run_id}")
def history_delete(run_id: str, sess=Depends(current)):
    hid = _hid(sess)
    if not history_store.delete(hid, run_id):
        raise ServiceError("Run not found in history.", 404)
    log.info("history delete %s", run_id)
    if (sess.get("batch") or {}).get("batch_id") == run_id:  # deleting the open batch -> fresh empty batch
        sess["batch"] = _new_batch()
        session_store.save(sess["sid"])
    elif sess["run"]["run_id"] == run_id:   # deleting the open run -> start a fresh empty run
        session_store.reset_run(sess["sid"])
    return {"ok": True, "items": history_store.list_items(hid), "run": _public_run(_run(sess))}




# --------------------------------------------------------------------------- #
# Batch mode - a separate workspace; the single run (sess["run"]) is never touched here.
# --------------------------------------------------------------------------- #
BATCH_MAX = 10
_ID_RE = re.compile(r"^[A-Z][A-Z0-9]+-\d+$")


def _new_batch() -> Dict[str, Any]:
    return {"kind": "batch", "batch_id": uuid.uuid4().hex[:12], "rid": "",
            "created": datetime.now().strftime("%Y-%m-%d %H:%M:%S"), "active": None, "items": [], "seq": 0}


def _batch(sess) -> Dict[str, Any]:
    if not sess.get("batch"):
        sess["batch"] = _new_batch()
    return sess["batch"]


def _item_summary(it) -> Dict[str, Any]:
    run = it["run"]
    cases = run.get("test_cases") or []
    up = run.get("upload") or {}
    acl = (run.get("ac_list") or _ac_list(run)) if cases else []
    return {k: it.get(k) for k in ("key", "source", "item_id", "title", "fetch", "ac_count", "gen", "ready",
                                    "target", "qtest", "filename", "template_ids")} | {
        "run_id": run["run_id"], "rid": run.get("rid", ""), "count": len(cases),
        "reviewed": sum(1 for c in cases if c.get("status") == "reviewed"), "dirty": bool(run.get("dirty")),
        "steps": sum(len(c.get("steps") or []) for c in cases),
        "uploaded": up.get("created", 0) if up else 0, "upload_total": up.get("total", 0) if up else 0,
        "cov_total": len(acl),
        "cov_uncovered": [x["id"] for x in acl if not any(x["id"] in (c.get("acs") or []) for c in cases)]}


def _public_batch(b) -> Dict[str, Any]:
    return {"batch_id": b["batch_id"], "rid": b.get("rid") or "", "created": b.get("created"),
            "active": b.get("active"), "max": BATCH_MAX, "template_ids": b.get("template_ids") or [], "items": [_item_summary(i) for i in b["items"]]}


def _bresp(sess, with_run=True, **extra):
    b = _batch(sess)
    out = {"ok": True, "mode": sess.get("mode") or "single", "batch": _public_batch(b),
           "single": {"rid": sess["run"].get("rid", ""), "story_id": sess["run"].get("story_id", ""),
                      "count": len(sess["run"].get("test_cases") or [])}}
    it = _bitem(b, b.get("active"))
    if it and with_run:
        out["run"] = _public_run(it["run"])
    out.update(extra)
    return out


def _need_item(sess, key) -> Dict[str, Any]:
    it = _bitem(_batch(sess), key)
    if not it:
        raise ServiceError("Batch item not found.", 404)
    return it


def _new_item(b, source, item_id="", title="") -> Dict[str, Any]:
    if len(b["items"]) >= BATCH_MAX:
        raise ServiceError(f"A batch holds at most {BATCH_MAX} items.")
    b["seq"] = b.get("seq", 0) + 1
    run = session_store._new_run()
    run.update(source="jira" if source == "jira" else "local", story_id=item_id)
    it = {"key": f"i{b['seq']}", "source": source, "item_id": item_id, "title": title,
          "fetch": {"status": "pending", "error": ""}, "ac_count": 0,
          "gen": {"status": "waiting", "error": "", "tokens": 0, "secs": 0}, "ready": False,
          "target": {"module_name": "", "module_id": "", "requirement_id": ""}, "qtest": {"status": "", "error": ""},
          "filename": "", "run": run}
    b["items"].append(it)
    return it


def _next_local_id(b) -> str:
    used = {i["item_id"] for i in b["items"]}
    n = 1
    while f"LOCAL-{n:03d}" in used:
        n += 1
    return f"LOCAL-{n:03d}"


@app.get("/api/batch")
def batch_get(sess=Depends(current)):
    return _bresp(sess, with_run=False)   # status refresh only; never swaps the open run in the browser


@app.post("/api/batch/enter")
def batch_enter(sess=Depends(current)):
    sess["mode"] = "batch"
    _batch(sess)
    _save(sess)
    log.info("batch mode on")
    return _bresp(sess)


@app.post("/api/batch/exit")
def batch_exit(sess=Depends(current)):
    sess["mode"] = "single"
    _save(sess)
    log.info("batch mode off")
    return {"ok": True, "mode": "single", "run": _public_run(sess["run"])}


@app.post("/api/batch/new")
def batch_new(sess=Depends(current)):
    sess["mode"] = "batch"
    _save(sess, history=True)   # the finished batch stays in history as one entry
    sess["batch"] = _new_batch()
    _save(sess)
    return _bresp(sess)


@app.post("/api/batch/items")
def batch_add(data: Dict[str, Any] = Body(...), sess=Depends(current)):
    b = _batch(sess)
    src = str(data.get("source", ""))
    added, errors = [], []
    if src == "jira":
        raw = data.get("ids") or []
        if isinstance(raw, str):
            raw = re.split(r"[\s,;]+", raw)
        for x in [str(r).strip().upper() for r in raw if str(r).strip()]:
            if not _ID_RE.match(x):
                errors.append(f"{x}: not a valid Jira ID")
            elif any(i["item_id"] == x for i in b["items"]):
                errors.append(f"{x}: already in this batch")
            elif len(b["items"]) >= BATCH_MAX:
                errors.append(f"{x}: batch is full ({BATCH_MAX} max)")
            else:
                added.append(_new_item(b, "jira", x)["key"])
    elif src == "local":
        text = str(data.get("text", "")).strip()
        if not text:
            raise ServiceError("Feature description is empty.")
        if len(text) > MAX_TEXT:
            raise ServiceError(f"Feature description is too long (max {MAX_TEXT:,} characters).")
        item_id = str(data.get("item_id", "")).strip().upper() or _next_local_id(b)
        if any(i["item_id"] == item_id for i in b["items"]):
            raise ServiceError(f"{item_id} is already in this batch.")
        title = str(data.get("title", "")).strip()[:120] or next((l.strip()[:90] for l in text.splitlines() if l.strip()), item_id)
        it = _new_item(b, "local", item_id, title)
        it["run"].update(feature_text=text, local_filename=str(data.get("filename", ""))[:200])
        acs = _ac_list(it["run"])
        it["run"]["ac_list"] = acs
        it.update(ac_count=len(acs), fetch={"status": "ok", "error": ""})
        it["gen"]["status"] = "queued"
        added.append(it["key"])
    elif src == "excel":
        fname = str(data.get("filename", ""))[:200]
        res = _parse_upload(data)
        if not res["cases"]:
            raise ServiceError(f"{fname}: no test cases found in the workbook.")
        m = re.search(r"[A-Za-z][A-Za-z0-9]+-\d+", fname)
        item_id = (str(data.get("item_id", "")).strip() or (m.group(0) if m else "")).upper() or _next_local_id(b)
        if any(i["item_id"] == item_id for i in b["items"]):
            raise ServiceError(f"{item_id} is already in this batch.")
        it = _new_item(b, "excel", item_id, f"Imported from {fname}")
        cases = parsing.sanitize_cases(excel_io.apply_mode([], res["cases"], "replace", res["needs_renumber"]))
        _validate_cases(cases)
        it["run"].update(test_cases=cases, imported=True, dirty=True, ac_list=res["ac_list"],
                         import_info={"filename": fname, "mode": "replace", "count": len(cases)})
        it.update(filename=fname, ac_count=len(res["ac_list"]), fetch={"status": "ok", "error": ""},
                  gen={"status": "imported", "error": "", "tokens": 0, "secs": 0})
        added.append(it["key"])
        _save(sess, history=True)
    else:
        raise ServiceError("Unknown batch source.")
    _save(sess)
    return _bresp(sess, added=added, errors=errors)


@app.post("/api/batch/items/{key}/fetch")
def batch_fetch(key: str, sess=Depends(current)):
    it = _need_item(sess, key)
    if it["source"] != "jira":
        raise ServiceError("Only Jira items are fetched.")
    try:
        info = services.fetch_jira(it["item_id"])
    except ServiceError as ex:
        it["fetch"] = {"status": "failed", "error": ex.message}
        it["gen"]["status"] = "waiting"
        _save(sess)
        return _bresp(sess)
    run = it["run"]
    run.update(source="jira", story_id=info["issue_key"], jira=info, feature_text=info["feature_text"])
    run["ac_list"] = _ac_list(run)
    it.update(title=info.get("summary") or it["item_id"], ac_count=len(run["ac_list"]))
    if info.get("usable") and info.get("feature_text"):
        it["fetch"] = {"status": "ok", "error": ""}
        if it["gen"]["status"] in ("waiting", "failed", "skipped", "cancelled"):
            it["gen"] = {"status": "queued", "error": "", "tokens": 0, "secs": 0}
    else:
        it["fetch"] = {"status": "failed", "error": info.get("message") or "Story has no usable acceptance criteria."}
    _save(sess)
    return _bresp(sess)


@app.patch("/api/batch/items/{key}")
def batch_patch(key: str, data: Dict[str, Any] = Body(...), sess=Depends(current)):
    it = _need_item(sess, key)
    if "ready" in data:
        if data["ready"] and not it["run"].get("test_cases"):
            raise ServiceError("This item has no test cases yet.")
        it["ready"] = bool(data["ready"])
    if "item_id" in data and it["source"] != "jira":
        new = str(data["item_id"]).strip().upper()
        if not new:
            raise ServiceError("ID cannot be empty.")
        if any(i["item_id"] == new and i is not it for i in _batch(sess)["items"]):
            raise ServiceError(f"{new} is already in this batch.")
        it["item_id"] = new
        it["run"]["story_id"] = new
    if isinstance(data.get("target"), dict):
        for k in ("module_name", "module_id", "requirement_id"):
            if k in data["target"]:
                it["target"][k] = str(data["target"][k] or "").strip()[:200]
    if "template_ids" in data:   # prompt templates: per-item override (null/[] = use batch selection)
        if it["source"] == "excel":
            raise ServiceError("Excel items are not generated, so they have no template.")
        it["template_ids"] = ptpl.resolve_ids(data["template_ids"]) if data["template_ids"] else None
    if "gen" in data and data["gen"] in ("queued", "skipped") and it["gen"]["status"] in ("queued", "skipped", "failed", "cancelled", "waiting"):
        if data["gen"] == "queued" and it["fetch"]["status"] != "ok":
            raise ServiceError("Fetch this item first.")
        it["gen"] = {"status": data["gen"], "error": "", "tokens": 0, "secs": 0}
    _save(sess, history=bool(it["run"].get("test_cases")))
    return _bresp(sess)


@app.delete("/api/batch/items/{key}")
def batch_remove(key: str, sess=Depends(current)):
    b = _batch(sess)
    it = _need_item(sess, key)
    if it["gen"]["status"] == "generating":
        raise ServiceError("Cancel generation before removing this item.")
    b["items"].remove(it)
    if b.get("active") == key:
        b["active"] = None
    _save(sess, history=True)
    return _bresp(sess)


@app.post("/api/batch/active")
def batch_active(data: Dict[str, Any] = Body(...), sess=Depends(current)):
    b = _batch(sess)
    key = data.get("key")
    if key and not _bitem(b, key):
        raise ServiceError("Batch item not found.", 404)
    b["active"] = key or None
    _save(sess)
    return _bresp(sess)


@app.post("/api/batch/items/{key}/generate")
def batch_generate(key: str, data: Dict[str, Any] = Body(default={}), sess=Depends(current)):
    """Generate one item. The frontend drives the queue sequentially (one request at a time)."""
    it = _need_item(sess, key)
    if it["fetch"]["status"] != "ok" or it["source"] == "excel":
        raise ServiceError("This item cannot be generated.")
    if it["gen"]["status"] == "generating":
        raise ServiceError("Already generating.")
    run = it["run"]
    model = str(data.get("model", "gemini")).lower()
    mode = "custom" if data.get("prompt_mode") == "custom" else "default"
    custom = str(data.get("custom_prompt", ""))[:50_000]
    run.update({"model": model, "prompt_mode": mode, "custom_prompt": custom})
    token = uuid.uuid4().hex
    it["gen"] = {"status": "generating", "error": "", "tokens": 0, "secs": 0, "token": token}
    _save(sess)
    t0 = time.time()
    try:
        tpl_ids = _tpl_ids_for(run, {"template_ids": it.get("template_ids") or _batch(sess).get("template_ids")}) if mode == "default" else None
        tpl_res = None
        if tpl_ids:
            tpl_res = _tpl_generate(run, model, tpl_ids)
            result, cases = tpl_res["result"], tpl_res["cases"]
        else:
            prompt, _ = services.build_prompt(mode, custom)
            result = services.generate(model, run["feature_text"], prompt)
            cases = parsing.markdown_to_test_cases(result["markdown"])
            if mode == "default" and cases:   # Default template implicitly used -> count it
                ptpl.mark_used([ptpl.DEFAULT_ID])
        if not cases:
            raise ServiceError("The model responded, but no test cases could be parsed from its output.", 422)
    except ServiceError as ex:
        if it["gen"].get("token") == token:
            it["gen"] = {"status": "failed", "error": ex.message, "tokens": 0, "secs": round(time.time() - t0, 1)}
            _save(sess)
        return _bresp(sess)
    if it["gen"].get("token") != token or _bitem(_batch(sess), key) is not it:   # cancelled / removed meanwhile
        return _bresp(sess, discarded=True)
    run["ac_list"] = _ac_list(run)
    parsing.suggest_acs(cases, run["ac_list"])
    cases = parsing.sanitize_cases(cases)
    run.update({"imported": False, "markdown": result["markdown"], "test_cases": cases, "dirty": True,
                "excel_path": "", "review": None, "upload": None, "rtm": [], "last_usage": result["usage"]})
    it["gen"] = {"status": "done", "error": "", "tokens": result["usage"].get("total_tokens") or 0,
                 "secs": round(time.time() - t0, 1)}
    it["ready"] = False
    _save(sess, history=True)
    for u in (tpl_res["usages"] if tpl_res else [result["usage"]]):
        _tok(sess, "Generate", u, run, item=it)
    return _bresp(sess)


@app.post("/api/batch/items/{key}/cancel")
def batch_cancel(key: str, sess=Depends(current)):
    it = _need_item(sess, key)
    if it["gen"]["status"] in ("generating", "queued"):
        it["gen"] = {"status": "cancelled", "error": "Cancelled", "tokens": 0, "secs": 0}
        _save(sess)
    return _bresp(sess)


@app.post("/api/batch/cancel_all")
def batch_cancel_all(sess=Depends(current)):
    for it in _batch(sess)["items"]:
        if it["gen"]["status"] in ("generating", "queued"):
            it["gen"] = {"status": "cancelled", "error": "Cancelled", "tokens": 0, "secs": 0}
    _save(sess)
    return _bresp(sess)


@app.post("/api/batch/items/{key}/qtest")
def batch_qtest(key: str, data: Dict[str, Any] = Body(default={}), sess=Depends(current)):
    it = _need_item(sess, key)
    run = it["run"]
    if not run.get("test_cases"):
        raise ServiceError("No test cases.")
    tgt = it["target"]
    msg = ""
    if it["source"] == "local" and not (tgt.get("module_id") or tgt.get("module_name")):
        raise ServiceError(f"{it['item_id']}: choose a module (name or ID) for local items.")
    req = tgt.get("requirement_id") or ""
    if not req and it["source"] in ("jira", "excel"):
        req = services.find_requirement(str(data.get("project_id", "")), it["item_id"]) or ""
        msg = (f"Requirement {req} found for {it['item_id']}" if req
               else f"No qTest requirement found for {it['item_id']} — uploaded without a link")
    if run.get("dirty") or not run.get("excel_path") or not Path(run["excel_path"]).exists():
        _write_excel(run)
    it["qtest"] = {"status": "uploading", "error": ""}
    try:
        result = services.qtest_upload(Path(run["excel_path"]), run["test_cases"], it["item_id"],
                                       str(data.get("project_id", "")), req, str(data.get("parent_module_id", "")),
                                       tgt.get("module_id", ""), tgt.get("module_name") or it["item_id"])
    except ServiceError as ex:
        it["qtest"] = {"status": "failed", "error": ex.message}
        _save(sess)
        return _bresp(sess)
    run["upload"] = result
    run["rtm"] = _build_rtm(run)
    full = result["created"] == result["total"]
    it["qtest"] = {"status": "done" if full else "partial", "error": "" if full else f"{result['total'] - result['created']} failed",
                   "message": msg, "requirement_id": req, "module_id": result.get("module_id")}
    _save(sess, history=True)
    return _bresp(sess)


@app.post("/api/batch/items/{key}/qtest/retry")
def batch_qtest_retry(key: str, sess=Depends(current)):
    it = _need_item(sess, key)
    run = it["run"]
    if not run.get("upload"):
        return batch_qtest(key, {}, sess)
    run["upload"] = services.qtest_retry_failed(run["upload"], run["test_cases"])
    run["rtm"] = _build_rtm(run)
    up = run["upload"]
    full = up["created"] == up["total"]
    it["qtest"].update(status="done" if full else "partial", error="" if full else f"{up['total'] - up['created']} failed")
    _save(sess, history=True)
    return _bresp(sess)


def _batch_zip(b, sess=None, hid=None):
    buf = io.BytesIO()
    n = 0
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for it in b["items"]:
            run = it["run"]
            if not run.get("test_cases"):
                continue
            if not run.get("excel_path") or run.get("dirty") or not Path(run["excel_path"]).exists():
                _write_excel(run)
            z.write(run["excel_path"], f"{it['item_id']}_{Path(run['excel_path']).name}")
            n += 1
    if not n:
        raise ServiceError("No item in this batch has test cases yet.")
    if sess is not None:
        session_store.save(sess["sid"])
    elif hid:   # opened from history without loading: keep the written Excel paths in the snapshot
        p = config.HISTORY_DIR / hid / "runs" / f"{b['batch_id']}.json"
        if p.exists():
            p.write_text(json.dumps(b, default=str), encoding="utf-8")
    name = f"{b.get('rid') or 'BATCH'}_test_cases.zip"
    return Response(buf.getvalue(), media_type="application/zip",
                    headers={"Content-Disposition": f'attachment; filename="{name}"'})


@app.get("/api/batch/excel.zip")
def batch_excel_zip(sess=Depends(current)):
    return _batch_zip(_batch(sess), sess)


# --------------------------------------------------------------------------- #
# Duplicate detection (similarity is computed locally in the browser; no AI call)
# --------------------------------------------------------------------------- #
@app.get("/api/settings/dup_threshold")
def get_dup_threshold(sess=Depends(current)):
    return {"ok": True, "value": settings_store.get_dup_threshold(), "min": settings_store.DUP_THRESHOLD_MIN,
            "max": settings_store.DUP_THRESHOLD_MAX, "default": settings_store.DUP_THRESHOLD_DEFAULT}


@app.put("/api/settings/dup_threshold")
def put_dup_threshold(data: Dict[str, Any] = Body(...), sess=Depends(current)):
    try:
        n = settings_store.set_dup_threshold(data.get("value"))
    except ValueError as e:
        raise ServiceError(str(e))
    log.info("duplicate threshold set to %d", n)
    return {"ok": True, "value": n}


@app.put("/api/review/dup_keep")
def put_dup_keep(data: Dict[str, Any] = Body(...), sess=Depends(current)):
    """'Keep both' decisions for this run: list of content-fingerprint pairs "a|b" (survive renumbering
    and Run history reopen). Stored on the run; the test cases themselves are not changed."""
    pairs = [str(p)[:80] for p in (data.get("pairs") or []) if isinstance(p, str) and "|" in p][:2000]
    run = _run(sess)
    run["dup_keep"] = sorted(set(pairs))
    _save(sess, history=bool(run.get("test_cases")))
    return {"ok": True, "dup_keep": run["dup_keep"]}


# --------------------------------------------------------------------------- #
# Compare runs (two runs of the same story from Run history; diff is computed in the browser)
# --------------------------------------------------------------------------- #
def _find_run(sess, run_id: str) -> Optional[Dict[str, Any]]:
    """The live run (single or open batch item) first, then history: single snapshots and batch children."""
    if sess["run"]["run_id"] == run_id:
        return sess["run"]
    for it in (sess.get("batch") or {}).get("items", []):
        if it["run"]["run_id"] == run_id:
            return it["run"]
    hid = sess.get("hid")
    snap = history_store.load_run(hid, run_id)
    if snap and snap.get("kind") != "batch":
        return snap
    for e in history_store.list_items(hid):
        if e.get("kind") == "batch" and any(c["run_id"] == run_id for c in e.get("children") or []):
            b = history_store.load_run(hid, e["run_id"]) or {}
            for it in b.get("items", []):
                if it["run"]["run_id"] == run_id:
                    return it["run"]
    return None


def _cmp_view(run, created) -> Dict[str, Any]:
    u = run.get("last_usage") or {}
    return {"run_id": run["run_id"], "rid": run.get("rid", ""), "story_id": run.get("story_id", ""),
            "created": created or run.get("created", ""), "test_cases": run.get("test_cases") or [],
            "ac_list": run.get("ac_list") or _ac_list(run), "tokens": u.get("total_tokens") or None,
            "imported": bool(run.get("imported") and not run.get("markdown"))}


@app.get("/api/compare")
def compare_runs(a: str, b: str, sess=Depends(current)):
    hid = _hid(sess)
    created = {e["run_id"]: e.get("created") for e in history_store.list_items(hid)}
    for e in history_store.list_items(hid):
        for c in e.get("children") or []:
            created[c["run_id"]] = e.get("created")
    ra, rb = _find_run(sess, a), _find_run(sess, b)
    if not ra or not rb:
        raise ServiceError("One of the runs is no longer in history.", 404)
    if (ra.get("story_id") or "").upper() != (rb.get("story_id") or "").upper() or not ra.get("story_id"):
        raise ServiceError("Pick 2 runs of the same story.")
    va, vb = _cmp_view(ra, created.get(a)), _cmp_view(rb, created.get(b))
    older, newer = sorted([va, vb], key=lambda r: (r["created"] or "", r["rid"] or ""))
    cur = _run(sess)["run_id"]
    log.info("compare %s vs %s", older["rid"], newer["rid"])
    return {"ok": True, "older": older, "newer": newer, "current": cur}


@app.post("/api/compare/excel")
def compare_excel(data: Dict[str, Any] = Body(...), sess=Depends(current)):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    rows = data.get("rows") or []
    wb = Workbook()
    ws = wb.active
    ws.title = "Compare"
    ws.append([str(data.get("title", ""))[:200]])
    ws.append(["Match", "Older TC", "Older name", "Current TC", "Current name", "Difference"])
    for c in ws[2]:
        c.font = Font(bold=True)
    fills = {"Added": "E6F6EC", "Removed": "FDEAEA", "Changed": "FFF4DC"}
    for r in rows[:5000]:
        vals = [str(r.get(k, ""))[:500] for k in ("match", "old_id", "old_name", "new_id", "new_name", "diff")]
        ws.append(vals)
        if vals[0] in fills:
            ws.cell(ws.max_row, 1).fill = PatternFill("solid", fgColor=fills[vals[0]])
    for col, w in zip("ABCDEF", (12, 10, 50, 10, 50, 50)):
        ws.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    name = re.sub(r"[^A-Za-z0-9_.-]", "_", str(data.get("filename") or "compare.xlsx"))[:80]
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{name}"'})


# --------------------------------------------------------------------------- #
# Token Usage page (additive)
# --------------------------------------------------------------------------- #
def _tq(request: Request):
    q = request.query_params
    return dict(d_from=q.get("from", "")[:10], d_to=q.get("to", "")[:10], mode=q.get("mode", ""), model=q.get("model", ""))


@app.get("/api/tokens")
def tokens(request: Request, sess=Depends(current)):
    return {"ok": True, **token_usage.query(**_tq(request))}


@app.get("/api/tokens/export.csv")
def tokens_csv(request: Request, sess=Depends(current)):
    body = token_usage.export_csv(**_tq(request))
    return Response(content=body, media_type="text/csv",
                    headers={"Content-Disposition": 'attachment; filename="token_usage.csv"'})


@app.get("/api/settings/tokens")
def get_token_settings(sess=Depends(current)):
    return {"ok": True, "settings": token_usage.get_settings()}


@app.put("/api/settings/tokens")
def put_token_settings(data: Dict[str, Any] = Body(...), sess=Depends(current)):
    try:
        return {"ok": True, "settings": token_usage.set_settings(data)}
    except ValueError as ex:
        raise ServiceError(str(ex))


# --------------------------------------------------------------------------- #
# Global search (header search / Ctrl+K) over the current run/batch + run history
# --------------------------------------------------------------------------- #
def _sr_hit(q, *vals):
    return any(q in str(v or "").lower() for v in vals)


@app.get("/api/search")
def global_search(q: str = "", limit: int = 8, sess=Depends(current)):
    q = (q or "").strip().lower()[:200]
    limit = max(1, min(int(limit or 8), 50))
    out = {"runs": [], "batches": [], "tickets": [], "test_cases": [], "acs": []}
    if not q:
        return {"ok": True, "q": q, **out}
    hid = sess.get("hid")
    entries = history_store.list_items(hid) if history_store.valid(hid) else []
    seen_t, seen_tc, seen_ac = set(), set(), set()

    def add(group, item, key=None, seen=None):
        if len(out[group]) >= limit or (seen is not None and key in seen):
            return
        if seen is not None:
            seen.add(key)
        out[group].append(item)

    def scan_run(run, ctx):
        """ctx: {run_id, rid, open_id (history id to open), key (batch item key)}"""
        if not run:
            return
        sid = run.get("story_id") or ""
        if sid and _sr_hit(q, sid, (run.get("jira") or {}).get("summary")):
            add("tickets", {**ctx, "story_id": sid, "title": (run.get("jira") or {}).get("summary") or ""}, sid, seen_t)
        for c in run.get("test_cases") or []:
            steps = " ".join(f"{s.get('description', '')} {s.get('expected', '')}" for s in c.get("steps") or [])
            if _sr_hit(q, c.get("id"), c.get("name"), c.get("description"), c.get("precondition"), steps):
                add("test_cases", {**ctx, "tc_id": c.get("id"), "name": c.get("name") or "", "story_id": sid},
                    (ctx["run_id"], c.get("id")), seen_tc)
        try:
            acl = run.get("ac_list") or _ac_list(run)
        except Exception:  # noqa: BLE001 - search must never fail on an odd run
            acl = []
        for a in acl or []:
            if _sr_hit(q, a.get("id"), a.get("title"), a.get("text")):
                add("acs", {**ctx, "ac_id": a.get("id"), "title": a.get("title") or a.get("text") or "", "story_id": sid},
                    (ctx["run_id"], a.get("id")), seen_ac)

    # current single run first (may not be in history yet)
    cur = sess.get("run") or {}
    scan_run(cur, {"run_id": cur.get("run_id"), "rid": cur.get("rid") or "", "open_id": cur.get("run_id"), "key": "", "current": True})
    for e in entries:
        if e.get("kind") == "batch":
            if _sr_hit(q, e.get("rid"), e.get("story_id"), e.get("title")) or any(
                    _sr_hit(q, c.get("rid"), c.get("story_id"), c.get("title")) for c in e.get("children") or []):
                add("batches", {"run_id": e["run_id"], "rid": e.get("rid"), "title": e.get("title"), "story_id": e.get("story_id"),
                                "count": e.get("count"), "created": e.get("created")})
            snap = history_store.load_run(hid, e["run_id"]) or {}
            for it in snap.get("items") or []:
                r = it.get("run") or {}
                if r.get("test_cases"):
                    scan_run(r, {"run_id": r.get("run_id"), "rid": r.get("rid") or "", "open_id": e["run_id"], "key": it.get("key"), "batch_rid": e.get("rid")})
            continue
        if _sr_hit(q, e.get("rid"), e.get("story_id"), e.get("title"), e.get("source")):
            add("runs", {"run_id": e["run_id"], "rid": e.get("rid"), "title": e.get("title"), "story_id": e.get("story_id"),
                         "count": e.get("count"), "status": e.get("status"), "created": e.get("created")})
        if e["run_id"] == cur.get("run_id"):
            continue
        scan_run(history_store.load_run(hid, e["run_id"]), {"run_id": e["run_id"], "rid": e.get("rid") or "", "open_id": e["run_id"], "key": ""})
    return {"ok": True, "q": q, **out}


@app.get("/api/health")
def health():
    return {"ok": True}


# --------------------------------------------------------------------------- #
# Prompt templates (additive). Interaction with the Generate step:
#  * Generate step "Custom prompt" -> always wins (templates ignored), exactly as before.
#  * Generate step "Default prompt" -> uses the template selection from the Input step (or batch/item
#    selection). Selection == [built-in Default] (the initial state) -> identical to the old Default path.
# --------------------------------------------------------------------------- #
def _tpl_ids_for(run, data) -> Optional[List[str]]:
    src = data.get("template_ids") if isinstance(data, dict) and data.get("template_ids") else run.get("template_ids")
    ids = ptpl.resolve_ids(src)
    run["template_ids"] = ids
    return None if ids == [ptpl.DEFAULT_ID] else ids


def _tpl_generate(run, model, ids) -> Dict[str, Any]:
    """Call the model once per template, tag each case with its source template, merge, renumber TC IDs."""
    ctx = ptpl.context_from_run(run)
    merged, usages, mds, prompts = [], [], [], []
    tot = {"model": model, "input_tokens": 0, "output_tokens": 0, "total_tokens": 0, "model_name": "", "templates": []}
    for tid in ids:
        t = ptpl.get(tid)
        if not t:
            continue
        prompt = ptpl.fill(t, ctx)
        prompts.append({"id": tid, "name": t["name"], "prompt": prompt})
        try:
            res = services.generate(model, run["feature_text"], prompt, tag="" if tid == ptpl.DEFAULT_ID else t["name"])
        except ServiceError as ex:
            raise ServiceError(f"Template '{t['name']}': {ex.message}", ex.status)
        cs = parsing.markdown_to_test_cases(res["markdown"])
        if not cs:
            raise ServiceError(f"Template '{t['name']}': the model responded, but no test cases could be parsed from its output.", 422)
        for c in cs:
            c["template"] = t["name"]
        merged += cs
        mds.append(res["markdown"])
        u = dict(res["usage"], template=t["name"])
        usages.append(u)
        for k in ("input_tokens", "output_tokens", "total_tokens"):
            tot[k] += int(u.get(k) or 0)
        tot["model_name"] = u.get("model_name") or tot["model_name"]
        tot["templates"].append(t["name"])
    for i, c in enumerate(merged):
        c["id"] = f"TC{i + 1:03d}"
    run["prompts_used"] = prompts
    ptpl.mark_used(ids)
    md = parsing.test_cases_to_markdown(merged) if merged else "\n\n".join(mds)
    return {"result": {"markdown": md, "usage": tot}, "cases": merged, "usages": usages}


@app.get("/api/templates")
def tpl_list(sess=Depends(current)):
    d = ptpl.list_all()
    for t in d["templates"]:
        if t["id"] == ptpl.DEFAULT_ID:
            try:
                t["text"] = ptpl.default_text()
            except ServiceError as ex:
                t["text"], t["error"] = "", ex.message
    return {"ok": True, **d, "placeholders": [{"name": n, "required": r} for n, r in ptpl.PLACEHOLDERS],
            "test_types": ptpl.TEST_TYPES, "detail_levels": ptpl.DETAIL}


@app.post("/api/templates")
def tpl_create(data: Dict[str, Any] = Body(...), sess=Depends(current)):
    return {"ok": True, **ptpl.create(data)}


@app.post("/api/templates/validate")
def tpl_validate(data: Dict[str, Any] = Body(...), sess=Depends(current)):
    return {"ok": True, **ptpl.validate(str(data.get("text") or ""))}


@app.post("/api/templates/preview")
def tpl_preview(data: Dict[str, Any] = Body(...), sess=Depends(current)):
    run = _run(sess)
    ctx = ptpl.context_from_run(run) if (run.get("feature_text") and data.get("use_run", True)) else {
        "story_title": "MTG-1427 · Refund API for partial orders",
        "description": "As a customer-support agent I want to call POST /orders/{id}/refunds with selected line items "
                       "so that a partial refund is issued to the original payment method.",
        "acceptance_criteria": "AC-1 Returns 201 with refund_id for valid items\nAC-2 Returns 422 when refund exceeds "
                               "refundable balance\nAC-3 Returns 401 without a valid bearer token"}
    t = {"id": "preview", "text": str(data.get("text") or ""), "options": data.get("options") or {}}
    return {"ok": True, "prompt": ptpl.fill(t, ctx), "values": ctx, **ptpl.validate(t["text"])}


@app.post("/api/templates/select")
def tpl_select(data: Dict[str, Any] = Body(...), sess=Depends(current)):
    run = _run(sess)
    run["template_ids"] = ptpl.resolve_ids(data.get("ids"))
    _save(sess)
    return {"ok": True, "template_ids": run["template_ids"]}


@app.get("/api/templates/selection")
def tpl_selection(item: str = "", sess=Depends(current)):
    """Current selection + the filled prompt(s) ('View prompt'). item=<batch key> for a batch item."""
    if item:
        it = _need_item(sess, item)
        run, src = it["run"], it.get("template_ids") or _batch(sess).get("template_ids")
    else:
        run = _run(sess)
        src = run.get("template_ids")
    ids = ptpl.resolve_ids(src)
    ctx = ptpl.context_from_run(run)
    prompts = [{"id": i, "name": ptpl.get(i)["name"], "prompt": ptpl.fill(ptpl.get(i), ctx)} for i in ids]
    custom = run.get("prompt_mode") == "custom"
    return {"ok": True, "template_ids": ids, "prompts": prompts, "custom_override": custom,
            "last_used": run.get("prompts_used") or []}


@app.put("/api/batch/templates")
def batch_templates(data: Dict[str, Any] = Body(...), sess=Depends(current)):
    b = _batch(sess)
    b["template_ids"] = ptpl.resolve_ids(data.get("ids"))
    _save(sess)
    return _bresp(sess)


@app.post("/api/templates/default/reset")
def tpl_reset_default(sess=Depends(current)):
    return {"ok": True, **ptpl.reset_default()}


@app.get("/api/templates/default/original")
def tpl_original_default(sess=Depends(current)):
    return {"ok": True, "text": services.read_original_prompt()}


@app.put("/api/templates/{tid}")
def tpl_update(tid: str, data: Dict[str, Any] = Body(...), sess=Depends(current)):
    return {"ok": True, **ptpl.update(tid, data)}


@app.delete("/api/templates/{tid}")
def tpl_delete(tid: str, sess=Depends(current)):
    ptpl.delete(tid)
    return {"ok": True}


@app.post("/api/templates/{tid}/duplicate")
def tpl_duplicate(tid: str, sess=Depends(current)):
    return {"ok": True, **ptpl.duplicate(tid)}


@app.post("/api/templates/{tid}/default")
def tpl_set_default(tid: str, sess=Depends(current)):
    ptpl.set_default(tid)
    return {"ok": True, "default_id": tid}


# --------------------------------------------------------------------------- #
# Review tools (additive): Export (xlsx / csv / qTest-import csv) + per-run Version history.
# Existing endpoints are untouched; versions are recorded by a small middleware after the
# existing save points succeed (generate / batch generate / review save / AI review accept / Excel import).
# --------------------------------------------------------------------------- #
from . import review_tools as rtools  # noqa: E402

_RT_VERSION_HOOKS = [
    ("POST", re.compile(r"^/api/generate$"), "generated"),
    ("POST", re.compile(r"^/api/batch/items/([^/]+)/generate$"), "generated"),
    ("PUT", re.compile(r"^/api/review/test_cases$"), "manual edits"),
    ("POST", re.compile(r"^/api/review/accept$"), "AI review applied"),
    ("POST", re.compile(r"^/api/excel/import$"), "Excel import"),
]


@app.middleware("http")
async def _rt_version_hook(request: Request, call_next):
    resp = await call_next(request)
    try:
        if resp.status_code == 200:
            for method, rx, note in _RT_VERSION_HOOKS:
                m = rx.match(request.url.path)
                if request.method == method and m:
                    sess = session_store.get(request.cookies.get(COOKIE))
                    if sess:
                        if m.groups():
                            it = _bitem(sess.get("batch"), m.group(1))
                            run = it["run"] if it else None
                        else:
                            run = _run(sess)
                        if run is not None:
                            rtools.record(run, note)
                    break
    except Exception:
        log.exception("version history hook failed (ignored)")
    return resp


def _rt_target_run(sess, item: str = "") -> Dict[str, Any]:
    if item:
        it = _bitem(sess.get("batch"), item)
        if not it:
            raise ServiceError("Batch item not found.", 404)
        return it["run"]
    return _run(sess)


@app.get("/api/export")
def rt_export(fmt: str = "xlsx", cols: str = "", steps: int = 1, reviewed: int = 0, item: str = "",
              sess=Depends(current)):
    """Download the run's (or a batch item's) test cases. Does not change run['excel_path'] / dirty."""
    run = _rt_target_run(sess, item)
    if fmt not in rtools.FORMATS:
        raise ServiceError("Unknown export format.")
    cases = parsing.sanitize_cases(run.get("test_cases"))
    cases = rtools.select_cases(cases, bool(reviewed))
    if not cases:
        raise ServiceError("No test cases to export" + (" (none are marked Reviewed)." if reviewed else "."))
    name = (run.get("story_id") or "LOCAL").replace("/", "_")
    if fmt == "qtest":
        data = rtools.csv_bytes(rtools.qtest_table(cases))
        return Response(data, media_type="text/csv; charset=utf-8",
                        headers={"Content-Disposition": f'attachment; filename="{name}_qtest_import.csv"'})
    clist = rtools.parse_cols(cols)
    if fmt == "csv":
        data = rtools.csv_bytes(rtools.table(cases, clist, bool(steps)))
        return Response(data, media_type="text/csv; charset=utf-8",
                        headers={"Content-Disposition": f'attachment; filename="{name}_testcases.csv"'})
    path = rtools.xlsx_file(run, cases, clist, bool(steps))
    return FileResponse(path, filename=f"{name}_testcases.xlsx",
                        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.get("/api/versions")
def rt_versions(sess=Depends(current)):
    run = _run(sess)
    vs = rtools.load(run)
    if not vs and run.get("test_cases"):   # runs created before version history existed
        rtools.record(run, "initial (first seen)")
        vs = rtools.load(run)
    return {"ok": True, "run_id": run.get("run_id"), "story_id": run.get("story_id") or "",
            "title": (run.get("jira") or {}).get("summary") or "", "versions": vs}


@app.post("/api/versions/{v}/restore")
def rt_restore(v: int, sess=Depends(current)):
    run = _run(sess)
    ver = next((x for x in rtools.load(run) if x.get("v") == v), None)
    if not ver:
        raise ServiceError(f"Version v{v} not found.", 404)
    cases = parsing.sanitize_cases(ver["cases"])
    _validate_cases(cases)
    run.update({"test_cases": cases, "dirty": True})
    _save(sess, history=True)
    nv = rtools.record(run, f"restored v{v}")
    return {"ok": True, "restored": v, "version": (nv or {}).get("v"), "run": _public_run(run)}


# --------------------------------------------------------------------------- #
# Record feature (additive): /api/rec/* — record browser actions into test cases (no AI).
# --------------------------------------------------------------------------- #
try:
    from . import recorder as _recorder  # noqa: E402
    _recorder.install(app, current, _run, _save, _public_run, _validate_cases, _RT_VERSION_HOOKS)
except Exception:  # never let the new feature break the existing app
    log.exception("Record feature failed to load (ignored)")
