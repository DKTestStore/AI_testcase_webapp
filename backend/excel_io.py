"""Excel export extras (review-only columns) and Excel import (upload) for the Review/Input steps.

The extra columns (AC Tags, Priority, Review Status) and the 'AC List' sheet are display/round-trip
only. qTest payloads are built by qtest_helpers.read_test_cases_from_excel, which reads only the six
standard columns, so these extras never reach qTest.
"""
import base64
import io
import re
from pathlib import Path
from typing import Any, Dict, List, Optional

import openpyxl
from openpyxl.styles import Font, PatternFill

from .parsing import PRIORITIES, STATUSES
from .services import ServiceError

MAX_BYTES = 5 * 1024 * 1024
STATUS_LABEL = {"not": "Not reviewed", "reviewed": "Reviewed", "changes": "Needs changes"}
_STATUS_IN = {"not reviewed": "not", "not": "not", "notreviewed": "not", "": "not", "pending": "not",
              "reviewed": "reviewed", "done": "reviewed", "approved": "reviewed",
              "needs changes": "changes", "needs change": "changes", "changes": "changes", "needschanges": "changes"}
EXTRA_HEADERS = ["AC Tags", "Priority", "Review Status"]
STD_HEADERS = ["Test Case ID", "Test Case Name", "Test Case Description", "Precondition",
               "Step Description", "Expected Result"]
# normalized header -> field
_HEAD = {
    "testcaseid": "id", "tcid": "id", "id": "id", "testcase": "id",
    "testcasename": "name", "name": "name", "title": "name",
    "testcasedescription": "description", "description": "description",
    "precondition": "precondition", "preconditions": "precondition",
    "stepdescription": "step", "step": "step", "steps": "step", "action": "step",
    "expectedresult": "expected", "expected": "expected", "expectedresults": "expected",
    "actags": "acs", "ac": "acs", "acceptancecriteria": "acs", "acs": "acs",
    "priority": "priority",
    "reviewstatus": "status", "status": "status",
}
REQUIRED = ["name", "step", "expected"]
LABELS = {"id": "Test Case ID", "name": "Test Case Name", "description": "Test Case Description",
          "precondition": "Precondition", "step": "Step Description", "expected": "Expected Result",
          "acs": "AC Tags", "priority": "Priority", "status": "Review Status"}


# --------------------------------------------------------------------------- export
def add_review_columns(path: Path, cases: List[Dict[str, Any]], ac_list: Optional[List[Dict[str, str]]] = None) -> None:
    """Append AC Tags / Priority / Review Status columns (first row of each TC) and an 'AC List' sheet."""
    wb = openpyxl.load_workbook(path)
    ws = wb.worksheets[0]
    col = ws.max_column + 1
    fill = PatternFill("solid", fgColor="E2F5E9")
    for k, h in enumerate(EXTRA_HEADERS):
        c = ws.cell(row=1, column=col + k, value=h)
        c.font = Font(bold=True)
        c.fill = fill
        ws.column_dimensions[openpyxl.utils.get_column_letter(col + k)].width = 16
    ti = -1
    for r in range(2, ws.max_row + 1):
        if str(ws.cell(row=r, column=1).value or "").strip():
            ti += 1
            if ti < len(cases):
                tc = cases[ti]
                ws.cell(row=r, column=col, value=", ".join(tc.get("acs") or []))
                ws.cell(row=r, column=col + 1, value=tc.get("priority") or "Medium")
                ws.cell(row=r, column=col + 2, value=STATUS_LABEL.get(tc.get("status") or "not", "Not reviewed"))
    if ac_list:
        al = wb.create_sheet("AC List")
        al.append(["AC ID", "Title"])
        for c in al[1]:
            c.font = Font(bold=True)
        for a in ac_list:
            al.append([a.get("id", ""), a.get("title", "")])
        al.column_dimensions["A"].width = 10
        al.column_dimensions["B"].width = 70
    wb.save(path)


def template_bytes() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Test Cases"
    ws.append(STD_HEADERS + EXTRA_HEADERS)
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.append(["TC001", "Example: Verify Generate Scenario Docs button", "Describe what is verified",
               "Required system state", "1. Click 'Generate Scenario Docs'", "Generate Scenario Documents page opens",
               "AC-01", "High", "Not reviewed"])
    ws.append(["", "", "", "", "2. Verify the page header", "Header reads 'Generate Scenario Documents'", "", "", ""])
    for col, w in zip("ABCDEFGHI", (12, 40, 40, 30, 40, 40, 14, 10, 16)):
        ws.column_dimensions[col].width = w
    al = wb.create_sheet("AC List")
    al.append(["AC ID", "Title"])
    al.append(["AC-01", "Example acceptance criterion title"])
    notes = wb.create_sheet("Instructions")
    for line in ["One row per step. Fill Test Case ID / Name / Description / Precondition only on the first step row.",
                 "Required columns: Test Case Name, Step Description, Expected Result (Test Case ID recommended).",
                 "Optional: AC Tags (comma separated, e.g. AC-01, AC-04), Priority (High/Medium/Low),",
                 "Review Status (Not reviewed / Reviewed / Needs changes). Optional columns are not sent to qTest."]:
        notes.append([line])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------- import
def _norm(h: Any) -> str:
    return re.sub(r"[^a-z]", "", str(h or "").lower())


def _s(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r"\s+", " ", str(v)).strip()


def decode_upload(filename: str, data_b64: str) -> bytes:
    if not filename.lower().endswith(".xlsx"):
        raise ServiceError("Only .xlsx files are supported (save the file as Excel Workbook .xlsx).")
    try:
        raw = base64.b64decode(data_b64.split(",")[-1], validate=False)
    except Exception:
        raise ServiceError("Could not read the uploaded file.")
    if len(raw) > MAX_BYTES:
        raise ServiceError("File is larger than 5 MB.")
    if not raw:
        raise ServiceError("The uploaded file is empty.")
    return raw


def parse_workbook(raw: bytes) -> Dict[str, Any]:
    try:
        wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception:
        raise ServiceError("The file is not a valid .xlsx workbook.")
    ws = wb.worksheets[0]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    # header = first non-empty row within the first 10
    hi = next((i for i, r in enumerate(rows[:10]) if any(_s(c) for c in r)), None)
    if hi is None:
        raise ServiceError("The first sheet is empty.")
    cols: Dict[str, int] = {}
    for j, h in enumerate(rows[hi]):
        f = _HEAD.get(_norm(h))
        if f and f not in cols:
            cols[f] = j
    missing = [LABELS[f] for f in REQUIRED if f not in cols]
    if missing:
        raise ServiceError("Missing required column(s): " + ", ".join(missing) +
                           ". Use 'Download template' for the expected format.")
    ac_titles: Dict[str, str] = {}
    for sh in wb.worksheets[1:]:
        if _norm(sh.title) == "aclist":
            for r in list(sh.iter_rows(values_only=True))[1:]:
                if r and _s(r[0]):
                    ac_titles[_s(r[0]).upper()] = _s(r[1]) if len(r) > 1 else ""
    get = lambda r, f: _s(r[cols[f]]) if f in cols and cols[f] < len(r) else ""
    cases: List[Dict[str, Any]] = []
    warnings: List[Dict[str, Any]] = []
    preview: List[Dict[str, Any]] = []
    for i, r in enumerate(rows[hi + 1:], start=hi + 2):
        vals = {f: get(r, f) for f in LABELS}
        if not any(vals.values()):
            continue
        last = cases[-1] if cases else None
        new_tc = (bool(vals["id"]) and (last is None or vals["id"] != last["id"])) or \
                 (not vals["id"] and bool(vals["name"]) and (last is None or vals["name"] != last["name"]))
        bad = []
        if new_tc or not cases:
            pr_raw = vals["priority"]
            pr = next((p for p in PRIORITIES if p.lower() == pr_raw.lower()), None)
            if pr_raw and not pr:
                warnings.append({"row": i, "tc": vals["id"], "field": "Priority",
                                 "message": f"Priority “{pr_raw}” is not valid (High / Medium / Low) — will import as Medium."})
                bad.append("priority")
            st_raw = vals["status"]
            st = _STATUS_IN.get(st_raw.lower().strip())
            if st_raw and not st:
                warnings.append({"row": i, "tc": vals["id"], "field": "Review Status",
                                 "message": f"Review Status “{st_raw}” is not valid (Not reviewed / Reviewed / Needs changes) — will import as Not reviewed."})
                bad.append("status")
            acs = [a.strip().upper().replace(" ", "") for a in re.split(r"[,;\n]+", vals["acs"]) if a.strip()]
            acs = [re.sub(r"^AC-?0*(\d+)$", lambda m: f"AC-{int(m.group(1)):02d}", a) for a in acs]
            if not vals["name"]:
                warnings.append({"row": i, "tc": vals["id"], "field": "Test Case Name", "message": "Test case has no name."})
                bad.append("name")
            cases.append({"id": vals["id"], "name": vals["name"], "description": vals["description"],
                          "precondition": vals["precondition"], "steps": [], "priority": pr or "Medium",
                          "status": st or "not", "acs": acs})
        if vals["step"] or vals["expected"]:
            cases[-1]["steps"].append({"description": vals["step"], "expected": vals["expected"]})
            if not vals["expected"] or not vals["step"]:
                warnings.append({"row": i, "tc": cases[-1]["id"], "field": "Step",
                                 "message": "Step has an empty " + ("Expected Result." if not vals["expected"] else "Step Description.")})
                bad.append("expected" if not vals["expected"] else "step")
        if len(preview) < 12 or bad:
            if len(preview) < 40:
                preview.append({"row": i, "values": vals, "bad": bad})
    if not cases:
        raise ServiceError("No test cases found in the first sheet.")
    for tc in cases:
        if not tc["steps"]:
            tc["steps"] = [{"description": "", "expected": ""}]
            warnings.append({"row": None, "tc": tc["id"], "field": "Steps", "message": f"{tc['id'] or tc['name']} has no steps."})
    ids = [c["id"] for c in cases]
    needs_renumber = any(not x for x in ids) or len(set(ids)) != len(ids)
    if needs_renumber:
        warnings.append({"row": None, "tc": "", "field": "Test Case ID",
                         "message": "Some Test Case IDs are blank or duplicated — IDs will be renumbered TC001…"})
    all_acs = sorted({a for c in cases for a in c["acs"]} | set(ac_titles), key=lambda a: (len(a), a))
    return {"cases": cases, "warnings": warnings, "preview": preview, "needs_renumber": needs_renumber,
            "columns": [LABELS[f] for f in LABELS if f in cols], "missing_optional": [LABELS[f] for f in LABELS if f not in cols],
            "ac_list": [{"id": a, "title": ac_titles.get(a, ""), "text": ""} for a in all_acs],
            "counts": {"test_cases": len(cases), "steps": sum(len(c["steps"]) for c in cases), "warnings": len(warnings),
                       "rows": len(rows) - hi - 1}}


def renumber(cases: List[Dict[str, Any]]) -> None:
    for i, c in enumerate(cases, 1):
        c["id"] = f"TC{i:03d}"


def apply_mode(current: List[Dict[str, Any]], imported: List[Dict[str, Any]], mode: str, needs_renumber: bool):
    if mode == "append":
        out = list(current) + imported
        renumber(out)
    elif mode == "merge":
        if needs_renumber:
            raise ServiceError("Merge by TC ID needs unique Test Case IDs in the file.")
        out = [dict(c) for c in current]
        idx = {c["id"]: k for k, c in enumerate(out)}
        for c in imported:
            if c["id"] in idx:
                out[idx[c["id"]]] = c
            else:
                out.append(c)
                idx[c["id"]] = len(out) - 1
        ids = [c["id"] for c in out]
        if len(set(ids)) != len(ids):
            renumber(out)
    else:
        out = imported
        if needs_renumber:
            renumber(out)
    return out
